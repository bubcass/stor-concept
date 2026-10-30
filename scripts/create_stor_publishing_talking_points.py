from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

OUT = "Stór publishing talking points.docx"

def shade(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    node = OxmlElement("w:shd")
    node.set(qn("w:fill"), fill)
    tc_pr.append(node)

def format_cell(cell, fill="FFFFFF"):
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
    shade(cell, fill)
    tc_pr = cell._tc.get_or_add_tcPr()
    borders = OxmlElement("w:tcBorders")
    for edge in ("top", "left", "bottom", "right"):
        node = OxmlElement(f"w:{edge}")
        node.set(qn("w:val"), "single")
        node.set(qn("w:sz"), "6")
        node.set(qn("w:color"), "D9D9D9")
        borders.append(node)
    tc_pr.append(borders)
    mar = OxmlElement("w:tcMar")
    for side in ("top", "start", "bottom", "end"):
        node = OxmlElement(f"w:{side}")
        node.set(qn("w:w"), "120" if side in ("top", "bottom") else "150")
        node.set(qn("w:type"), "dxa")
        mar.append(node)
    tc_pr.append(mar)

def cell_text(cell, text, bold=False, color="000000", size=9.5):
    cell.text = ""
    p = cell.paragraphs[0]
    p.paragraph_format.space_after = Pt(0)
    p.paragraph_format.line_spacing = 1.08
    r = p.add_run(text)
    r.bold = bold
    r.font.name = "Aptos"
    r.font.size = Pt(size)
    r.font.color.rgb = RGBColor.from_string(color)

def bullet(doc, text):
    p = doc.add_paragraph(style="List Bullet")
    p.paragraph_format.space_after = Pt(2.5)
    p.paragraph_format.line_spacing = 1.05
    p.paragraph_format.left_indent = Inches(0.18)
    p.paragraph_format.first_line_indent = Inches(-0.13)
    r = p.add_run(text)
    r.font.name = "Aptos"
    r.font.size = Pt(10)

doc = Document()
section = doc.sections[0]
section.page_width = Inches(8.5)
section.page_height = Inches(11)
section.top_margin = Inches(0.55)
section.bottom_margin = Inches(0.55)
section.left_margin = Inches(0.65)
section.right_margin = Inches(0.65)
doc.styles["Normal"].font.name = "Aptos"
doc.styles["Normal"].font.size = Pt(10.5)
title_style_ppr = doc.styles["Title"]._element.get_or_add_pPr()
for border in title_style_ppr.findall(qn("w:pBdr")):
    title_style_ppr.remove(border)

title = doc.add_paragraph(style="Title")
title.paragraph_format.space_after = Pt(2)
title_ppr = title._p.get_or_add_pPr()
for border in title_ppr.findall(qn("w:pBdr")):
    title_ppr.remove(border)
r = title.add_run("Stór publishing talking points")
r.font.name = "Aptos Display"
r.font.size = Pt(25)
r.font.color.rgb = RGBColor(0, 0, 0)
sub = doc.add_paragraph()
sub.paragraph_format.space_after = Pt(12)
r = sub.add_run("Why we use Mammoth for import and structured SvelteKit rendering for publication")
r.italic = True
r.font.size = Pt(11)

p = doc.add_paragraph()
p.paragraph_format.space_after = Pt(10)
p.paragraph_format.line_spacing = 1.15
r = p.add_run("Core message. ")
r.bold = True
r.font.size = Pt(11)
r = p.add_run("We can publish Mammoth HTML directly, but that would make the Word document our publishing system. Stór uses Mammoth to import content, then renders recognised structures as a consistent, accessible web product.")
r.font.size = Pt(11)

def heading(text):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(8)
    p.paragraph_format.space_after = Pt(4)
    r = p.add_run(text)
    r.bold = True
    r.font.size = Pt(13)

heading("The distinction to make")
table = doc.add_table(rows=1, cols=2)
table.alignment = WD_TABLE_ALIGNMENT.CENTER
table.autofit = False
for cell, text in zip(table.rows[0].cells, ("Publish imported Mammoth HTML", "Transform and render in Stór")):
    cell.width = Inches(3.48)
    format_cell(cell, "173F5F")
    cell_text(cell, text, True, "FFFFFF", 10)

comparisons = [
    ("Generic headings, lists and tables", "A designed, accessible report experience"),
    ("A membership list remains a membership list", "It becomes linked member cards with portraits and fallbacks"),
    ("Word styles and document quirks reach the site", "Stór controls design, responsiveness and accessibility"),
    ("Corrections may require changes to source files", "Improve one component and every matching report benefits"),
    ("Difficult to add consistent rich web features", "Reusable blocks for video, maps, charts, tables and resources"),
]
for index, pair in enumerate(comparisons):
    cells = table.add_row().cells
    for cell, text in zip(cells, pair):
        cell.width = Inches(3.48)
        format_cell(cell, "F4F7F9" if index % 2 == 0 else "FFFFFF")
        cell_text(cell, text)

heading("Useful responses in the room")
bullet(doc, "“Mammoth is not being replaced. It is the import layer that gives us useful document structure.”")
bullet(doc, "“The question is not HTML versus Svelte. Svelte also delivers HTML; it gives us reusable, tested components around the imported content.”")
bullet(doc, "“A linked bulleted list can become an accessible committee-member directory with official portraits. That is hard to achieve reliably by publishing generic HTML alone.”")
bullet(doc, "“If our goal were simply to place Word documents online quickly, direct HTML publication would be valid. Stór is intended to publish reports as a durable public web service.”")
bullet(doc, "“This is progressive: ordinary prose stays ordinary prose. We only upgrade recognised patterns where a native web component adds clear value.”")

heading("A short closing line")
p = doc.add_paragraph()
p.paragraph_format.space_after = Pt(0)
p.paragraph_format.line_spacing = 1.1
r = p.add_run("Word remains the authoring format. Mammoth remains the bridge into the web. Stór provides the consistent public representation.")
r.bold = True
r.font.size = Pt(10.5)

doc.save(OUT)
print(OUT)
