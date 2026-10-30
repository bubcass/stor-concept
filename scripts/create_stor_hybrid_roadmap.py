from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

OUT = "Stór hybrid publishing roadmap.docx"

def border_and_fill(cell, fill):
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
    tcpr = cell._tc.get_or_add_tcPr()
    shade = OxmlElement("w:shd")
    shade.set(qn("w:fill"), fill)
    tcpr.append(shade)
    borders = OxmlElement("w:tcBorders")
    for side in ("top", "left", "bottom", "right"):
        edge = OxmlElement(f"w:{side}")
        edge.set(qn("w:val"), "single")
        edge.set(qn("w:sz"), "6")
        edge.set(qn("w:color"), "D9D9D9")
        borders.append(edge)
    tcpr.append(borders)
    margins = OxmlElement("w:tcMar")
    for side, size in (("top", "125"), ("start", "150"), ("bottom", "125"), ("end", "150")):
        edge = OxmlElement(f"w:{side}")
        edge.set(qn("w:w"), size)
        edge.set(qn("w:type"), "dxa")
        margins.append(edge)
    tcpr.append(margins)

def write(cell, text, *, bold=False, colour="000000", size=9.6):
    cell.text = ""
    p = cell.paragraphs[0]
    p.paragraph_format.space_after = Pt(0)
    p.paragraph_format.line_spacing = 1.08
    r = p.add_run(text)
    r.bold = bold
    r.font.name = "Aptos"
    r.font.size = Pt(size)
    r.font.color.rgb = RGBColor.from_string(colour)

doc = Document()
section = doc.sections[0]
section.top_margin = Inches(0.55)
section.bottom_margin = Inches(0.55)
section.left_margin = Inches(0.65)
section.right_margin = Inches(0.65)
doc.styles["Normal"].font.name = "Aptos"
doc.styles["Normal"].font.size = Pt(10.5)
title_style_ppr = doc.styles["Title"]._element.get_or_add_pPr()
for border in title_style_ppr.findall(qn("w:pBdr")):
    title_style_ppr.remove(border)

title = doc.add_paragraph(style="Title")
title.paragraph_format.space_after = Pt(2)
r = title.add_run("Stór hybrid publishing roadmap")
r.font.name = "Aptos Display"
r.font.size = Pt(25)
r.font.color.rgb = RGBColor(0, 0, 0)
sub = doc.add_paragraph()
sub.paragraph_format.space_after = Pt(12)
r = sub.add_run("A phased route from Word import to a governed publishing service and refined public experience")
r.italic = True
r.font.size = Pt(11)

p = doc.add_paragraph()
p.paragraph_format.space_after = Pt(10)
p.paragraph_format.line_spacing = 1.15
r = p.add_run("Aim. ")
r.bold = True
r.font.size = Pt(11)
r = p.add_run("Keep Word as the authoring format, build a reusable structured-report pipeline, then connect it to the existing Laravel and Silverstripe estate before progressively delivering the richer Svelte public interface.")
r.font.size = Pt(11)

table = doc.add_table(rows=1, cols=3)
table.alignment = WD_TABLE_ALIGNMENT.CENTER
table.autofit = False
widths = (Inches(1.15), Inches(2.85), Inches(2.95))
for cell, text, width in zip(table.rows[0].cells, ("Phase", "Primary work", "Outcome and decision gate"), widths):
    cell.width = width
    border_and_fill(cell, "173F5F")
    write(cell, text, bold=True, colour="FFFFFF", size=10)

rows = [
    (
        "1\nStructured import",
        "Build and test the Mammoth.js import pipeline. Convert DOCX to a framework-neutral report model; recognise headings, paragraphs, tables, links, images and known patterns such as committee membership. Preserve the original source file and capture conversion warnings.",
        "A repeatable, testable import produces structured report JSON and native blocks. Confirm that representative reports—including committee reports—convert reliably before platform integration.",
    ),
    (
        "2\nGoverned publishing",
        "Integrate the import service with Laravel, which owns uploads, versions, roles, review, publishing state, audit and APIs. Laravel serves structured report data to Silverstripe for the existing web estate, or directly to the initial public presentation layer.",
        "A controlled system of record: drafts can be reviewed and approved, published reports are API-addressable, and Silverstripe can surface them without duplicating content. Agree the report/block API contract here.",
    ),
    (
        "3\nPublic experience",
        "Introduce the Svelte frontend progressively against the stable API. Start with report templates, contents navigation, responsive typography and print; then add components such as member cards, charts, maps, video, saved reports and richer search.",
        "A high-quality public experience without changing the editorial system. Svelte is an implementation layer, not the system of record; ordinary report content continues to work throughout.",
    ),
]
for index, row in enumerate(rows):
    cells = table.add_row().cells
    for cell, text, width in zip(cells, row, widths):
        cell.width = width
        border_and_fill(cell, "F4F7F9" if index % 2 == 0 else "FFFFFF")
        write(cell, text, bold=(cell is cells[0]), size=9.2)

heading = doc.add_paragraph()
heading.paragraph_format.space_before = Pt(12)
heading.paragraph_format.space_after = Pt(4)
r = heading.add_run("Responsibilities by layer")
r.bold = True
r.font.size = Pt(13)

points = [
    ("Mammoth and transformation service", "Understand the DOCX and produce a portable structured report model."),
    ("Laravel", "Own the authoritative records, workflow, permissions, versioning and publishing APIs."),
    ("Silverstripe", "Continue to provide existing web-estate integration and editorial context while the service is introduced."),
    ("Svelte", "Render the public report experience and specialised interactive blocks from the published API contract."),
]
for label, body in points:
    p = doc.add_paragraph(style="List Bullet")
    p.paragraph_format.space_after = Pt(2.5)
    p.paragraph_format.line_spacing = 1.06
    p.paragraph_format.left_indent = Inches(0.18)
    p.paragraph_format.first_line_indent = Inches(-0.13)
    r = p.add_run(label + ". ")
    r.bold = True
    r.font.size = Pt(10)
    r = p.add_run(body)
    r.font.size = Pt(10)

p = doc.add_paragraph()
p.paragraph_format.space_before = Pt(7)
p.paragraph_format.space_after = Pt(0)
p.paragraph_format.line_spacing = 1.1
r = p.add_run("Guiding principle. ")
r.bold = True
r.font.size = Pt(10.5)
r = p.add_run("Invest first in a framework-neutral content contract. That keeps the Word import, the backend service and the eventual frontend independently replaceable.")
r.font.size = Pt(10.5)

doc.save(OUT)
print(OUT)
